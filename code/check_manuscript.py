# -*- coding: utf-8 -*-
"""Structural and submission hygiene checker for the v51 ReliFuse manuscript."""

import argparse
import io
import os
import re
import sys
from pathlib import Path

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
DEFAULT_MAIN = ROOT / "paper" / "v51_ReliFuse_paper.md"
DEFAULT_SUPP = ROOT / "paper" / "v51_ReliFuse_Supplementary.md"
DEFAULT_DOCX = [ROOT / "paper" / "v51_ReliFuse_paper.docx",
                ROOT / "paper" / "v51_ReliFuse_Supplementary.docx"]


def roman_to_int(value):
    table = {"I": 1, "V": 5, "X": 10}
    total = prev = 0
    for ch in reversed(value):
        num = table[ch]
        total = total - num if num < prev else total + num
        prev = max(prev, num)
    return total


class Report:
    def __init__(self):
        self.rows = []

    def add(self, level, name, detail=""):
        self.rows.append((level, name, detail))

    def ok(self, name, detail=""):
        self.add("OK", name, detail)

    def warn(self, name, detail=""):
        self.add("WARN", name, detail)

    def fail(self, name, detail=""):
        self.add("FAIL", name, detail)

    def dump(self):
        width = max([len(x[1]) for x in self.rows] + [12])
        fail = warn = 0
        for level, name, detail in self.rows:
            print(f"[{level:<4}] {name:<{width}}  {detail}")
            fail += level == "FAIL"
            warn += level == "WARN"
        print("-" * 78)
        print(f"FAIL={fail}  WARN={warn}  OK={len(self.rows)-fail-warn}")
        return fail, warn


def read(path):
    return Path(path).read_text(encoding="utf-8")


def check_sequence(values, label, rep, start=1):
    if not values:
        rep.fail(f"{label} defined", "none found")
        return
    if len(values) != len(set(values)):
        duplicates = sorted({v for v in values if values.count(v) > 1})
        rep.fail(f"{label} unique", f"duplicates: {duplicates}")
    else:
        rep.ok(f"{label} unique", f"{len(values)} objects")
    expected = list(range(start, start + len(values)))
    if values != expected:
        rep.fail(f"{label} order", f"actual {values}, expected {expected}")
    else:
        rep.ok(f"{label} order", f"{start}..{max(values)}")


def check_images(main_text, supp_text, rep):
    base = DEFAULT_MAIN.parent
    all_images = re.findall(r"^!\[[^\]]*\]\(([^)]+)\)", main_text + "\n" + supp_text, re.M)
    missing = []
    for rel in all_images:
        if not (base / rel).resolve().exists():
            missing.append(rel)
    if missing:
        rep.fail("image paths", "; ".join(missing))
    else:
        rep.ok("image paths", f"{len(all_images)} image links resolve")


def check_main_objects(main_text, rep):
    lines = main_text.splitlines()
    table_caps = []
    table_cap_lines = set()
    for i, line in enumerate(lines):
        m = re.match(r"^Table ([IVX]+)\.\s", line)
        if m:
            table_caps.append(roman_to_int(m.group(1)))
            table_cap_lines.add(i)
    nums = table_caps
    check_sequence(nums, "main tables", rep)

    table_refs = set()
    for i, line in enumerate(lines):
        if i in table_cap_lines:
            continue
        for m in re.finditer(r"Table ([IVX]+)(?![A-Za-z])", line):
            table_refs.add(roman_to_int(m.group(1)))
    unreferenced = [x for x in nums if x not in table_refs]
    nonexistent = sorted(table_refs - set(nums))
    if unreferenced or nonexistent:
        rep.fail("main table references", f"unreferenced={unreferenced}; nonexistent={nonexistent}")
    else:
        rep.ok("main table references", f"{len(nums)} captions referenced")

    fig_caps = []
    fig_cap_lines = set()
    for i, line in enumerate(lines):
        if not line.startswith("!["):
            continue
        m = re.search(r"^!\[Fig\.\s*(\d+)\.", line)
        if m:
            fig_caps.append(int(m.group(1)))
            fig_cap_lines.add(i)
    check_sequence(fig_caps, "main figures", rep)
    fig_refs = set()
    for i, line in enumerate(lines):
        if i in fig_cap_lines:
            continue
        for m in re.finditer(r"Fig\.\s*(\d+)", line):
            fig_refs.add(int(m.group(1)))
    unreferenced = [x for x in fig_caps if x not in fig_refs]
    nonexistent = sorted(fig_refs - set(fig_caps))
    if unreferenced or nonexistent:
        rep.fail("main figure references", f"unreferenced={unreferenced}; nonexistent={nonexistent}")
    else:
        rep.ok("main figure references", f"{len(fig_caps)} captions referenced")


def check_supplement_objects(main_text, supp_text, rep):
    supp_tables = [int(x) for x in re.findall(r"^##\s*补充表\s*(\d+)\.", supp_text, re.M)]
    supp_figs = [int(x) for x in re.findall(r"^###\s*补充图\s*(\d+)\.", supp_text, re.M)]
    check_sequence(supp_tables, "supplement tables", rep, start=0)
    check_sequence(supp_figs, "supplement figures", rep)
    text = main_text + "\n" + supp_text
    t_refs = set(int(x) for x in re.findall(r"补充表\s*(\d+)", text))
    f_refs = set(int(x) for x in re.findall(r"补充图\s*(\d+)", text))
    t_missing = sorted(t_refs - set(supp_tables))
    f_missing = sorted(f_refs - set(supp_figs))
    if t_missing or f_missing:
        rep.fail("supplement cross-references", f"missing tables={t_missing}; missing figures={f_missing}")
    else:
        rep.ok("supplement cross-references", f"tables 0..{max(supp_tables)}; figures 1..{max(supp_figs)}")


def check_bibliography(main_text, rep):
    if "\n## References" not in main_text:
        rep.fail("bibliography section", "not found")
        return
    body, refs = main_text.split("\n## References", 1)
    used = set(int(x) for x in re.findall(r"\[(\d+)\]", body))
    defined = set(int(x) for x in re.findall(r"^\[(\d+)\]", refs, re.M))
    missing = sorted(used - defined)
    unused = sorted(defined - used)
    if missing or unused or defined != set(range(1, len(defined) + 1)):
        rep.fail("bibliography", f"missing={missing}; unused={unused}; count={len(defined)}")
    else:
        rep.ok("bibliography", f"{len(defined)} numbered references used")


def check_source_hygiene(main_text, supp_text, rep):
    combined = main_text + "\n" + supp_text
    forbidden = ["CalFuse", "CLF", "MKMP-CLF", "TODO:", "TBD:", "44_../"]
    hits = {token: combined.count(token) for token in forbidden if combined.count(token)}
    if hits:
        rep.fail("forbidden source tokens", str(hits))
    else:
        rep.ok("forbidden source tokens", "none")
    for label, text in [("main", main_text), ("supplement", supp_text)]:
        if text.count("**") % 2 or text.count("`") % 2:
            rep.warn(f"{label} Markdown delimiters", "odd number of ** or ` delimiters")
        else:
            rep.ok(f"{label} Markdown delimiters", "balanced")


def check_abstract_and_contributions(main_text, max_abstract_words, rep):
    m = re.search(r"^## Abstract\s*\n(.*?)(?=^## )", main_text, re.M | re.S)
    if not m:
        rep.fail("abstract", "section not found")
        return
    words = re.findall(r"\b[\w’'-]+\b", m.group(1))
    if len(words) > max_abstract_words:
        rep.fail("abstract length", f"{len(words)} words > {max_abstract_words}")
    else:
        rep.ok("abstract length", f"{len(words)} words")
    contributions = re.findall(r"^\d+\.\s+", main_text, re.M)
    if len(contributions) != 4:
        rep.warn("contribution count", f"found {len(contributions)}, expected 4")
    else:
        rep.ok("contribution count", "4")


def docx_text(path):
    from docx import Document
    doc = Document(path)
    chunks = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            chunks.extend(cell.text for cell in row.cells)
    return "\n".join(chunks), len(doc.paragraphs), len(doc.tables)


def check_docx(paths, rep):
    for path in paths:
        if not path.exists():
            rep.fail("DOCX exists", f"{path.name}: missing")
            continue
        try:
            text, paragraphs, tables = docx_text(path)
        except Exception as exc:
            rep.fail("DOCX readable", f"{path.name}: {exc}")
            continue
        forbidden = ["CalFuse", "CLF", "MKMP-CLF", "TODO:", "TBD:", "[[EQ", "missing figure", "44_../"]
        hits = {x: text.count(x) for x in forbidden if text.count(x)}
        if hits:
            rep.fail("DOCX tokens", f"{path.name}: {hits}")
        else:
            rep.ok("DOCX tokens", f"{path.name}: clean")
        if "**" in text or "`" in text:
            rep.fail("DOCX Markdown residue", f"{path.name}: raw ** or ` present")
        else:
            rep.ok("DOCX Markdown residue", f"{path.name}: none")
        rep.ok("DOCX structure", f"{path.name}: {paragraphs} paragraphs, {tables} tables")


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--main", default=str(DEFAULT_MAIN))
    parser.add_argument("--supp", default=str(DEFAULT_SUPP))
    parser.add_argument("--abstract-max-words", type=int, default=300)
    parser.add_argument("--strict", action="store_true")
    args = parser.parse_args()

    main_text = read(args.main)
    supp_text = read(args.supp)
    rep = Report()
    check_images(main_text, supp_text, rep)
    check_main_objects(main_text, rep)
    check_supplement_objects(main_text, supp_text, rep)
    check_bibliography(main_text, rep)
    check_abstract_and_contributions(main_text, args.abstract_max_words, rep)
    check_source_hygiene(main_text, supp_text, rep)
    check_docx(DEFAULT_DOCX, rep)

    print(f"# manuscript check: {Path(args.main).name}")
    fails, warns = rep.dump()
    return 1 if fails or (args.strict and warns) else 0


if __name__ == "__main__":
    sys.exit(main())
