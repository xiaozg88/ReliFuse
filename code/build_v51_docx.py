# -*- coding: utf-8 -*-
"""Build the ReliFuse v51 manuscript DOCX from Markdown."""

import os
import re
import sys

from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import parse_xml

HERE = os.path.normpath(os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
# omml_helpers.py is vendored next to this script (copied from the upstream
# v1 pipeline) so that the release repository is self-contained.
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from omml_helpers import EQ  # noqa: E402

MD = os.path.join(HERE, "paper", "v51_ReliFuse_paper.md")
OUT = os.path.join(HERE, "paper", "v51_ReliFuse_paper.docx")
SUPP_MD = os.path.join(HERE, "paper", "v51_ReliFuse_Supplementary.md")
SUPP_OUT = os.path.join(HERE, "paper", "v51_ReliFuse_Supplementary.docx")

IMG_RE = re.compile(r"^!\[(.*)\]\((.*)\)$")
EQ_RE = re.compile(r"^\[\[EQ:.*?:(\d+)\]\]$")
INLINE_RE = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)")


def set_font(obj, name="SimSun", size=None, bold=None, color=None, italic=None):
    if hasattr(obj, "font"):
        f = obj.font
        f.name = name
        if size is not None:
            f.size = Pt(size)
        if bold is not None:
            f.bold = bold
        if italic is not None:
            f.italic = italic
        if color is not None:
            f.color.rgb = RGBColor(*color)
    if hasattr(obj, "element"):
        rpr = obj.element.get_or_add_rPr()
        rfonts = rpr.find(qn("w:rFonts"))
        if rfonts is None:
            rfonts = rpr.makeelement(qn("w:rFonts"), {})
            rpr.append(rfonts)
        rfonts.set(qn("w:eastAsia"), name)


def add_formatted_runs(paragraph, text, size=10.5, bold=False, italic=False, color=None, font="SimSun"):
    """Render the small Markdown subset used by the manuscript sources."""
    for part in INLINE_RE.split(text):
        if not part:
            continue
        token_bold = part.startswith("**") and part.endswith("**") and len(part) > 4
        token_italic = (not token_bold and part.startswith("*") and part.endswith("*") and len(part) > 2)
        token_code = part.startswith("`") and part.endswith("`") and len(part) > 2
        content = part[2:-2] if token_bold else part[1:-1] if token_italic or token_code else part
        run = paragraph.add_run(content)
        if token_code:
            set_font(run, "Consolas", size=size, bold=bold, italic=italic, color=color)
        else:
            set_font(run, font, size=size, bold=bold or token_bold, italic=italic or token_italic, color=color)


def set_cell(cell, text, bold=False, size=8.5):
    cell.text = ""
    p = cell.paragraphs[0]
    add_formatted_runs(p, text, size=size, bold=bold)
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(1)


def add_table(doc, header, rows):
    table = doc.add_table(rows=1, cols=len(header))
    table.style = "Table Grid"
    for j, h in enumerate(header):
        set_cell(table.rows[0].cells[j], h, bold=True, size=8.5)
    for row in rows:
        cells = table.add_row().cells
        for j, v in enumerate(row):
            if j < len(cells):
                set_cell(cells[j], v, size=8.5)


def add_image(doc, caption, path, base=None):
    if not os.path.isabs(path):
        cand = os.path.normpath(os.path.join(base or HERE, path))
        if os.path.exists(cand):
            path = cand
    if not os.path.exists(path):
        p = doc.add_paragraph()
        run = p.add_run("(missing figure: " + path + ")")
        set_font(run, "SimSun", 9, italic=True, color=(0xAA, 0x33, 0x33))
        return
    try:
        doc.add_picture(path, width=Inches(6.0))
    except Exception:
        pass
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(caption)
    set_font(run, "SimSun", 9, bold=True)


def build(md_path, out_path, title_size=16):
    doc = Document()
    normal = doc.styles["Normal"]
    set_font(normal, "SimSun", 10.5)
    normal.paragraph_format.space_after = Pt(4)
    for sec in doc.sections:
        sec.top_margin = Inches(0.9)
        sec.bottom_margin = Inches(0.9)
        sec.left_margin = Inches(1.0)
        sec.right_margin = Inches(1.0)

    lines = open(md_path, encoding="utf-8").read().splitlines()
    i, seen_h1 = 0, 0
    while i < len(lines):
        line = lines[i].rstrip()
        if not line:
            i += 1
            continue
        em = EQ_RE.match(line.strip())
        if em:
            p = doc.add_paragraph()
            p.paragraph_format.tab_stops.add_tab_stop(Inches(3.25), WD_TAB_ALIGNMENT.CENTER)
            p.paragraph_format.tab_stops.add_tab_stop(Inches(6.5), WD_TAB_ALIGNMENT.RIGHT)
            p.add_run("\t")
            key = int(em.group(1))
            if key in EQ:
                p._p.append(parse_xml(EQ[key]))
            else:
                run = p.add_run(f"(eq {key})")
                set_font(run, "SimSun", 10.5)
            p.add_run("\t(%d)" % key)
            i += 1
            continue
        img = IMG_RE.match(line.strip())
        if img:
            add_image(doc, img.group(1), img.group(2), base=os.path.dirname(md_path))
            i += 1
            continue
        if line.startswith("|"):
            tbl = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                tbl.append(lines[i].strip())
                i += 1
            rows = []
            for r in tbl:
                cells = [c.strip() for c in r.strip("|").split("|")]
                if all(set(c) <= {"-", ":"} for c in cells):
                    continue
                rows.append(cells)
            if rows:
                add_table(doc, rows[0], rows[1:])
            continue
        if line.startswith("# "):
            p = doc.add_paragraph()
            run = p.add_run(line[2:])
            if seen_h1 == 0:
                set_font(run, "SimSun", title_size, bold=True, color=(0, 0, 0))
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                set_font(run, "SimSun", 12, bold=True, color=(0x44, 0x44, 0x44), italic=True)
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            seen_h1 += 1
            i += 1
            continue
        if line.startswith("## "):
            p = doc.add_paragraph(style="Heading 2")
            run = p.add_run(line[3:])
            set_font(run, "SimSun", 12, bold=True, color=(0x1F, 0x3B, 0x64))
            i += 1
            continue
        if line.startswith("### "):
            p = doc.add_paragraph(style="Heading 3")
            run = p.add_run(line[4:])
            set_font(run, "SimSun", 11, bold=True, color=(0x33, 0x33, 0x33))
            i += 1
            continue
        if line.startswith("- ") or line.startswith("* "):
            p = doc.add_paragraph(style="List Bullet")
            add_formatted_runs(p, line[2:], size=10.5)
            i += 1
            continue
        if line.startswith("> "):
            p = doc.add_paragraph()
            add_formatted_runs(p, line[2:], size=10, italic=True, color=(0x55, 0x55, 0x55))
            i += 1
            continue
        p = doc.add_paragraph()
        add_formatted_runs(p, line, size=10.5)
        i += 1

    doc.save(out_path)
    print("saved", out_path)


if __name__ == "__main__":
    build(MD, OUT, title_size=16)
    build(SUPP_MD, SUPP_OUT, title_size=14)
