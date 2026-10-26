# -*- coding: utf-8 -*-
"""Audit the English v51 supplement and its cross-references to the main text."""
import re
import sys
from pathlib import Path

try:
    sys.stdout.reconfigure(encoding="utf-8")
except Exception:
    pass

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
MAIN = ROOT / "paper" / "v51_ReliFuse_paper.md"
SUPP = ROOT / "paper" / "v51_ReliFuse_Supplementary_EN.md"
DOCX = ROOT / "paper" / "v51_ReliFuse_Supplementary_EN.docx"

m = MAIN.read_text(encoding="utf-8")
s = SUPP.read_text(encoding="utf-8")
fail = 0

def check(name, ok, detail=""):
    global fail
    print(f"[{'OK' if ok else 'FAIL':<4}] {name:<28} {detail}")
    fail += not ok

tabs = [int(x) for x in re.findall(r"^## Supplementary Table (\d+)\.", s, re.M)]
figs = [int(x) for x in re.findall(r"^### Supplementary Fig\. (\d+)\.", s, re.M)]
check("table sequence", tabs == list(range(max(tabs) + 1)), f"0..{max(tabs) if tabs else -1}")
check("figure sequence", figs == list(range(1, 10)), f"1..{max(figs) if figs else 0}")
main_tab_refs = set(int(x) for x in re.findall(r"Supplementary Table\s*(\d+)", m))
main_fig_refs = set(int(x) for x in re.findall(r"Supplementary Fig\.\s*(\d+)", m))
check("main table targets", main_tab_refs <= set(tabs), f"referenced {sorted(main_tab_refs)}")
check("main figure targets", main_fig_refs <= set(figs), f"referenced {sorted(main_fig_refs)}")
supp_main_refs = set(re.findall(r"main-text Tables? ([IVX]+)", s)) | set(re.findall(r"main-text Table ([IVX]+)", s))
valid_roman = {"I","II","III","IV","V","VI","VII","VIII","IX","X","XI","XII","XIII","XIV","XV"}
check("main-text table refs", supp_main_refs <= valid_roman, f"referenced {sorted(supp_main_refs)}")
links = re.findall(r"^!\[[^\]]*\]\(([^)]+)\)", s, re.M)
missing = [x for x in links if not (SUPP.parent / x).resolve().exists()]
check("image links", not missing, f"{len(links)} links; missing={missing}")
cjk = sum("\u4e00" <= c <= "\u9fff" for c in s)
check("English-only source", cjk == 0, f"CJK chars={cjk}")
check("forbidden tokens", not any(x in s for x in ["CalFuse", "CLF", "TODO:", "TBD:", "[[EQ"]), "legacy names/placeholders")

from docx import Document
d = Document(DOCX)
text = "\n".join(p.text for p in d.paragraphs)
for table in d.tables:
    for row in table.rows:
        text += "\n" + "\n".join(cell.text for cell in row.cells)
images = sum(1 for rel in d.part.rels.values() if "image" in rel.reltype)
check("DOCX structure", len(d.tables) == len(tabs) and images == len(links), f"{len(d.paragraphs)} paragraphs; {len(d.tables)} tables; {images} images")
cjk_docx = sum("\u4e00" <= c <= "\u9fff" for c in text)
check("DOCX English-only", cjk_docx == 0, f"CJK chars={cjk_docx}")
check("DOCX formatting", "**" not in text and "`" not in text and "[[EQ" not in text, "no Markdown residue/placeholders")

print("-" * 78)
print(f"FAIL={fail}")
sys.exit(1 if fail else 0)
